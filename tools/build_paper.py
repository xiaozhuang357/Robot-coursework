"""从课程论文 markdown 正文生成符合《智能机器人软件设计方法与实践》格式规范的 Word 文档。

格式规则的唯一来源是教师附件《封面及格式规范》：
A4、四边 2.5cm、全文 1.5 倍行距、正文首行缩进 2 字符、一级标题另起一页、
封面无页码且正文从第 1 页编号，标题黑体、正文宋体小四、图表题宋体五号、
英文数字 Times New Roman、图表按章编号、代码块 Consolas 五号不缩进。
"""

from pathlib import Path
from zipfile import ZipFile
import textwrap
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
VERSION = "V2.0"
SOURCE = ROOT / f"LineGuard_课程论文_{VERSION}.md"
DEST = ROOT / f"LineGuard_课程论文_{VERSION}.docx"
ASSETS = ROOT / "paper" / "assets"
TEMPLATE = ROOT / "《智能机器人软件设计方法与实践》封面及格式规范.docx"

EA_BODY, EA_HEAD, EA_CODE = "宋体", "黑体", "宋体"
WEST, WEST_CODE = "Times New Roman", "Consolas"
SZ_BODY, SZ_H1, SZ_H2, SZ_H3 = 12, 15, 14, 12
SZ_CAPTION, SZ_CODE, SZ_REF = 10.5, 10.5, 10.5
LINE = 1.5
PAGE_W, PAGE_H, MARGIN = 21.0, 29.7, 2.5
USABLE_W_CM = PAGE_W - 2 * MARGIN
MAX_FIG_W_IN = USABLE_W_CM / 2.54
MAX_FIG_H_IN = 2.65


# ---------- 基础排版工具 ----------

def set_font(run, ea=EA_BODY, size=SZ_BODY, bold=False, western=WEST):
    """同时设置西文与中文字体：西文与数字走 Times New Roman / Consolas，中文走黑体或宋体。"""
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "cs"):
        fonts.set(qn(f"w:{slot}"), western)
    fonts.set(qn("w:eastAsia"), ea)


def no_space(paragraph):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)


def set_indent_chars(paragraph, chars=200):
    """用 firstLineChars 表达"首行缩进 2 字符"，随字号自适应，而不是固定磅值。"""
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(int(chars / 100 * SZ_BODY * 20)))


def body_paragraph(doc, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    paragraph = doc.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.line_spacing = LINE
    no_space(paragraph)
    if indent:
        set_indent_chars(paragraph)
    return paragraph


def inline(paragraph, text, size=SZ_BODY, ea=EA_BODY, bold=False):
    """渲染行内内容：`代码` 走 Consolas，**强调** 走加粗，其余走正文。"""
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            set_font(paragraph.add_run(part[1:-1]), EA_CODE, size, bold, WEST_CODE)
        elif part.startswith("**") and part.endswith("**"):
            set_font(paragraph.add_run(part[2:-2]), ea, size, True)
        else:
            set_font(paragraph.add_run(part), ea, size, bold)


def page_break_before(paragraph):
    paragraph.paragraph_format.page_break_before = True


def add_page_field(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "1"
    run.append(value)
    field.append(run)
    paragraph._p.append(field)


def restart_page_numbering(section):
    """正文节从第 1 页开始编号：封面不占编号。"""
    sect_pr = section._sectPr
    pg_num = OxmlElement("w:pgNumType")
    pg_num.set(qn("w:start"), "1")
    sect_pr.append(pg_num)


def setup_section(section):
    section.page_width, section.page_height = Cm(PAGE_W), Cm(PAGE_H)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Cm(MARGIN))
    section.header_distance = section.footer_distance = Cm(1.5)
    # 模板的隐式文档网格会把 1.5 倍行距再次吸附放大，以明文规范为准关闭网格。
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)


def png_size(path: Path):
    header = path.read_bytes()[:24]
    return int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")


# ---------- 文档构件 ----------

def insert_figure(doc, png: Path, caption: str):
    width_px, height_px = png_size(png)
    width = min(MAX_FIG_W_IN, MAX_FIG_H_IN / (height_px / width_px))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(png), width=Inches(width))
    for node in paragraph._p.xpath(".//wp:docPr"):
        node.set("descr", caption)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = LINE
    cap.paragraph_format.space_before = Pt(0)
    cap.paragraph_format.space_after = Pt(6)
    set_font(cap.add_run(caption), EA_BODY, SZ_CAPTION)


def add_table(doc, rows, caption):
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.line_spacing = LINE
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(2)
    cap.paragraph_format.keep_with_next = True
    set_font(cap.add_run(caption), EA_BODY, SZ_CAPTION)

    count = len(rows[0])
    table = doc.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    width = Cm(USABLE_W_CM)
    ratios = [0.22, 0.43, 0.35]
    if caption.startswith("表 3-"):
        ratios = [0.30, 0.32, 0.38]
    if caption.startswith("表 5-"):
        ratios = [0.30, 0.58, 0.12]
    if count != 3:
        ratios = [1 / count] * count
    for col, ratio in zip(table.columns, ratios):
        col.width = int(width * ratio)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    props.append(borders)
    for index, cells in enumerate(rows):
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(cells):
            cell = row.cells[j]
            cell.width = int(width * ratios[j])
            margins = OxmlElement("w:tcMar")
            for side, amount in (("top", 60), ("bottom", 60), ("left", 80), ("right", 80)):
                node = OxmlElement(f"w:{side}")
                node.set(qn("w:w"), str(amount))
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            cell._tc.get_or_add_tcPr().append(margins)
            if index == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "EEF2F4")
                cell._tc.get_or_add_tcPr().append(shade)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = LINE
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, value, SZ_BODY, EA_BODY, index == 0)
    tail = doc.add_paragraph()
    tail.paragraph_format.line_spacing = 1.0
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.space_after = Pt(8)
    set_font(tail.add_run(""), EA_BODY, SZ_CAPTION)


def add_code(doc, description, code_lines):
    before = doc.add_paragraph()
    before.paragraph_format.line_spacing = LINE
    no_space(before)
    set_font(before.add_run(""), EA_BODY, SZ_CODE)
    if description:
        note = doc.add_paragraph(style="Code Caption")
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        note.paragraph_format.line_spacing = LINE
        note.paragraph_format.space_before = Pt(10)
        note.paragraph_format.space_after = Pt(2)
        note.paragraph_format.keep_with_next = True
        set_font(note.add_run(description), EA_BODY, SZ_CAPTION)
    for line in code_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = LINE
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        set_font(p.add_run(line if line.strip() else "\u00a0"), EA_CODE, SZ_CODE, False, WEST_CODE)
    after = doc.add_paragraph()
    after.paragraph_format.line_spacing = LINE
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(0)
    set_font(after.add_run(""), EA_BODY, SZ_CAPTION)


# ---------- 正文解析 ----------

def heading_level(line):
    match = re.match(r"^(#{1,3})\s+(.*)$", line)
    if not match:
        return None, None
    return len(match[1]), match[2].strip()


def chapter_label(text):
    numbered = re.match(r"^(\d+)\s", text)
    if numbered:
        return numbered[1]
    if text.startswith("附录"):
        return "A"
    return None


def build_body(doc, content):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    setup_section(section)
    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.line_spacing = 1.0
    add_page_field(footer)
    restart_page_numbering(section)

    lines = content.splitlines()
    chapter = None
    figures = tables = 0
    total_figures = total_tables = 0
    first_heading = True
    pending_figure = pending_table = pending_code = None
    is_abstract = False
    index = 0
    while index < len(lines):
        raw = lines[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line.startswith("<!--") and line.endswith("-->"):
            marker = line[4:-3].strip()
            key, _, value = marker.partition("：")
            if key == "图注":
                pending_figure = value.strip()
            elif key == "图源":
                figures += 1
                total_figures += 1
                numbering = f"{chapter}-{figures}" if chapter else str(figures)
                caption = f"图 {numbering}  {pending_figure or value.strip()}"
                png = ASSETS / value.strip()
                if not png.is_file():
                    raise SystemExit(f"缺少配图 {png}")
                insert_figure(doc, png, caption)
                pending_figure = None
            elif key == "表注":
                pending_table = value.strip()
            elif key == "代码":
                pending_code = value.strip()
            index += 1
            continue

        level, text = heading_level(line)
        if level:
            if level == 1:
                is_abstract = text.startswith("摘要")
                chapter = chapter_label(text)
                figures = tables = 0
                display = re.sub(r"^(\d+)\s+", r"第 \1 章 ", text)
                p = doc.add_paragraph(display, style="Heading 1")
                no_space(p)
                p.paragraph_format.line_spacing = LINE
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(4)
                if not first_heading:
                    page_break_before(p)
                first_heading = False
            else:
                p = doc.add_paragraph(text, style=f"Heading {level}")
                no_space(p)
                p.paragraph_format.line_spacing = LINE
                p.paragraph_format.space_before = Pt(7 if level == 2 else 5)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
            index += 1
            continue

        if line.startswith("```"):
            block = []
            index += 1
            while index < len(lines) and lines[index].strip() != "```":
                block.append(lines[index])
                index += 1
            add_code(doc, pending_code, block)
            pending_code = None
            index += 1
            continue

        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [item.strip() for item in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", item) for item in cells):
                    rows.append(cells)
                index += 1
            tables += 1
            total_tables += 1
            numbering = f"{chapter}-{tables}" if chapter else str(tables)
            add_table(doc, rows, f"表 {numbering}  {pending_table or '数据表'}")
            pending_table = None
            continue

        no_indent = is_abstract or line.startswith("关键词") or line.startswith("[")
        align = WD_ALIGN_PARAGRAPH.LEFT if line.startswith("[") else WD_ALIGN_PARAGRAPH.JUSTIFY
        p = body_paragraph(doc, indent=not no_indent, align=align)
        inline(p, line, SZ_REF if line.startswith("[") else SZ_BODY)
        index += 1

    return total_figures, total_tables


def prepare_reference_cover(doc, title):
    """以教师原文件的封面为基础，保留校徽、表格与学期，仅填内容槽。"""
    remove = False
    for child in list(doc.element.body):
        if child.tag == qn("w:sectPr"):
            continue
        if "格式规范（本页删除）" in "".join(child.itertext()):
            remove = True
        if remove:
            doc.element.body.remove(child)
    setup_section(doc.sections[0])
    for p in doc.paragraphs:
        if p.text.strip() == "题目自拟":
            p.clear()
            p.style = doc.styles["Title"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = LINE
            no_space(p)
            set_font(p.add_run(title), EA_HEAD, 22, True)
        elif p.text.strip() and not p._p.xpath(".//w:drawing") and not p._p.xpath(".//w:pict"):
            p.paragraph_format.line_spacing = LINE
            # 模板正文的中西文字体按其文字规范统一，保留原有字号及对齐。
            for r in p.runs:
                size = r.font.size.pt if r.font.size else 16
                set_font(r, EA_HEAD if "课程" in p.text else EA_BODY, size, bool(r.bold))
    values = ("2024303524", "张立烁", title, "")
    table = doc.tables[0]
    for row, value in zip(table.rows, values):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.line_spacing = LINE
        p = row.cells[1].paragraphs[0]
        p.clear()
        p.paragraph_format.line_spacing = LINE
        no_space(p)
        set_font(p.add_run(value), EA_BODY, 12 if value == title else 16)
    # 双行题目占用原空白槽的高度，移除表格后的三行空段，保证封面仍为一页。
    after_table = False
    removed = 0
    for child in list(doc.element.body):
        if child.tag == qn("w:tbl"):
            after_table = True
        elif after_table and child.tag == qn("w:p") and not "".join(child.itertext()).strip() and removed < 3:
            doc.element.body.remove(child)
            removed += 1
    return doc


def inject_source_excerpts(content):
    """源码片段自动抽取，防止论文与附件各改一份后产生漂移。"""
    base = ROOT / "lineguard-system/backend/src/main/java/com/lineguard/service"
    state_source = (base / "TaskStateMachine.java").read_text()
    state = re.search(r"    public void check\(.*?\n    }", state_source, re.S)[0]
    command_source = (base / "MissionStateService.java").read_text()
    command = re.search(r"        if \(commandRepository.existsByTaskIdAndCommandId.*?\n        }", command_source, re.S)[0]
    state = textwrap.dedent(state)
    command = textwrap.dedent(command)
    command = command.replace("task.getId(), commandId))", "task.getId(),\n        commandId))")
    snippets = {"state_check": ("TaskStateMachine 状态校验源码", state),
                "command_dedup": ("MissionStateService 指令去重源码", command)}
    for key, (caption, code) in snippets.items():
        content = content.replace(f"<!-- 源码：{key} -->", f"<!-- 代码：{caption} -->\n```java\n{code}\n```")
    return content


def main():
    content = SOURCE.read_text(encoding="utf-8")
    title_cn = re.search(r"<!--\s*题目：(.+?)\s*-->", content)[1].strip()
    body_text = inject_source_excerpts(re.sub(r"<!--\s*(题目|英文题目)：.*?-->", "", content))

    doc = Document(TEMPLATE)
    for name in ("Title", "Caption", "Code Caption", "Heading 1", "Heading 2", "Heading 3"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(SZ_BODY)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = LINE
    normal.paragraph_format.first_line_indent = Pt(0)
    normal.paragraph_format.keep_with_next = False
    normal.paragraph_format.widow_control = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "false")
    normal._element.get_or_add_pPr().append(snap)
    for name in ("Normal", "Title"):
        style = doc.styles[name]
        style.font.name = WEST
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = name == "Title"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_BODY)
    for level, size in ((1, SZ_H1), (2, SZ_H2), (3, SZ_H3)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = WEST
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        for slot in ("ascii", "hAnsi", "cs"):
            fonts.set(qn(f"w:{slot}"), WEST)
        fonts.set(qn("w:eastAsia"), EA_HEAD)

    prepare_reference_cover(doc, title_cn)
    figures, tables = build_body(doc, body_text)

    for root in (doc.styles.element, doc.element):
        for border in list(root.xpath(".//w:pBdr")):
            border.getparent().remove(border)

    doc.core_properties.title = title_cn
    doc.core_properties.author = "张立烁"
    doc.core_properties.comments = ""
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    temporary = DEST.with_suffix(".building.docx")
    doc.save(temporary)
    # 保留无关的原始模板包部件，避免 python-docx 重序列化改变校徽和主题等资源。
    preserved = ("word/media/image1.png", "word/theme/theme1.xml", "word/fontTable.xml", "docProps/custom.xml")
    with ZipFile(TEMPLATE) as reference, ZipFile(temporary) as built, ZipFile(DEST, "w") as final:
        for entry in built.infolist():
            final.writestr(entry, reference.read(entry.filename) if entry.filename in preserved else built.read(entry.filename))
    temporary.unlink()

    han = len(re.findall(r"[\u4e00-\u9fff]", body_text))
    print(f"DOCX: {DEST.name}")
    print(f"题目: {title_cn}")
    print(f"正文汉字: {han}; 插图: {figures}; 表格: {tables}")


if __name__ == "__main__":
    main()
