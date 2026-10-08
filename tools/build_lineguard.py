"""从计划书正文生成可编辑 Word 文档，并把正文中的 mermaid 图源渲染为配图。"""

from pathlib import Path
import re
import subprocess

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT


ROOT = Path(__file__).resolve().parents[1]
VERSION = "V2.0"
SOURCE = ROOT / f"LineGuard_项目计划书_{VERSION}.md"
DEST = ROOT / f"LineGuard_项目计划书_{VERSION}.docx"
ASSETS = ROOT / ".lineguard_qa" / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)
MERMAID_THEME = ROOT / "tools" / "mermaid-theme.json"
SCREENSHOTS = ROOT / ".lineguard_qa" / "screenshots"
BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"
def set_font(run, name=BODY_FONT, size=11, bold=False, color="000000"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor.from_string(color)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), name)


def link(paragraph, label, url=None, anchor=None, size=11):
    node = OxmlElement("w:hyperlink")
    if url:
        node.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    else:
        node.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "24546B" if url else "000000")
    props.append(color)
    fonts = OxmlElement("w:rFonts")
    for slot in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{slot}"), BODY_FONT)
    props.append(fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    props.append(sz)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = label
    run.append(text_node)
    node.append(run)
    paragraph._p.append(node)


def inline(paragraph, text, size=11, bold=False, color="000000"):
    pattern = r"(`[^`]+`|\[[^\]]+\]\(https?://[^)]+\))"
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, "Consolas", size - 1, bold, color)
        elif part.startswith("[") and "](" in part:
            matched = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            link(paragraph, matched[1], matched[2], size=size)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size, bold=bold, color=color)


def bookmark(paragraph, identifier, name):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def png_size(path: Path) -> tuple[int, int]:
    header = path.read_bytes()[:24]
    return int.from_bytes(header[16:20], "big"), int.from_bytes(header[20:24], "big")


def render_mermaid(source: str, slug: str) -> Path:
    """把正文里的 mermaid 图源渲染为 PNG，.mmd 与 PNG 同名同源。"""
    mmd = ASSETS / f"{slug}.mmd"
    png = ASSETS / f"{slug}.png"
    mmd.write_text(source.rstrip() + "\n", encoding="utf-8")
    result = subprocess.run(
        ["mmdc", "-i", str(mmd), "-o", str(png), "-c", str(MERMAID_THEME),
         "-b", "white", "-s", "2"],
        capture_output=True, text=True,
    )
    if result.returncode != 0 or not png.is_file():
        raise SystemExit(f"mermaid 渲染失败（{slug}）：{result.stderr.strip() or result.stdout.strip()}")
    return png


def insert_figure(doc, png: Path, caption: str, alt: str):
    width_px, height_px = png_size(png)
    max_width, max_height = 6.4, 7.1
    width = min(max_width, max_height / (height_px / width_px))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.add_run().add_picture(str(png), width=Inches(width))
    for shape in paragraph._p.xpath(".//wp:docPr"):
        shape.set("descr", alt)
    caption_paragraph = doc.add_paragraph(caption, style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(10)


def add_table(doc, rows):
    count = len(rows[0])
    proportions = {
        3: [0.19, 0.40, 0.41],
        4: [0.16, 0.28, 0.28, 0.28],
        5: [0.28, 0.15, 0.07, 0.16, 0.34],
        8: [0.15, 0.10, 0.105, 0.105, 0.115, 0.095, 0.115, 0.215],
        2: [0.24, 0.76],
    }[count]
    # 按字段长度进一步调整，避免说明列挤占短字段。
    if rows[0][0] in ("触发情况", "检查项", "风险"):
        proportions = [0.21, 0.39, 0.40]
    table = doc.add_table(rows=0, cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    total = 6.8
    for col, fraction in zip(table.columns, proportions):
        col.width = Inches(total * fraction)
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "5")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in [("top", 95), ("bottom", 95), ("left", 90), ("right", 90)]:
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    props.append(margins)
    for index, cells in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            trpr.append(OxmlElement("w:tblHeader"))
        for j, value in enumerate(cells):
            cell = row.cells[j]
            cell.width = Inches(total * proportions[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "DCE8EF" if index == 0 else ("F6F8FA" if index % 2 == 0 else "FFFFFF"))
            cell._tc.get_or_add_tcPr().append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.18
            p.paragraph_format.keep_with_next = index <= 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 or (count == 8 and j < 7) else WD_ALIGN_PARAGRAPH.LEFT
            inline(p, value, size=9.5 if count == 8 else 10, bold=index == 0)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.line_spacing = Pt(4)
    set_font(after.add_run(""), size=3)


def main():
    content = SOURCE.read_text(encoding="utf-8")
    # 先清理旧配图，避免图数变化后留下与正文不对应的孤儿文件。
    for stale in list(ASSETS.glob("figure-*.png")) + list(ASSETS.glob("figure-*.mmd")):
        stale.unlink()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.85)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.7)
    sec.header_distance = sec.footer_distance = Inches(0.3)
    sec.different_first_page_header_footer = True
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3", "Caption"):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.italic = False
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for level, size in [(1, 17), (2, 13), (3, 11.5)]:
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(15 if level == 1 else 11)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    doc.core_properties.title = f"LineGuard 项目计划书 {VERSION}"
    doc.core_properties.subject = "偏远难达地区架空输电线路常驻式智能运维机器人系统"
    doc.core_properties.author = "LineGuard 项目组"
    doc.core_properties.keywords = "工业机器人,输电运维,OPGW,Java,智能巡检"
    doc.core_properties.comments = ""
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run(f"LineGuard  项目计划书 {VERSION}"), HEADING_FONT, 8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("第 "), size=9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    set_font(footer.add_run(" 页"), size=9)

    # 封面保留正式标题和版本背景。
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_before = Pt(78)
    p.paragraph_format.space_after = Pt(22)
    set_font(p.add_run("LineGuard"), HEADING_FONT, 36, True)
    p = doc.add_paragraph(style="Title")
    p.paragraph_format.space_after = Pt(22)
    set_font(p.add_run(f"项目计划书 {VERSION}"), HEADING_FONT, 25, True)
    p = doc.add_paragraph(style="Subtitle")
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(30)
    set_font(p.add_run("偏远难达地区架空输电线路\n常驻式智能运维机器人系统"), HEADING_FONT, 18)
    for text in ["项目定位  架空输电线路常驻式近距巡检系统", "产品阶段  单档距核心闭环验证", "版本日期  2026 年 9 月 10 日"]:
        p = doc.add_paragraph()
        inline(p, text, 11)
    intro = re.search(r"版本日期：[^\n]*\n\n(.*?)(?=\n\n# 1)", content, re.S)[1]
    for index, text in enumerate(intro.split("\n\n")):
        p = doc.add_paragraph()
        if index == 0:
            p.paragraph_format.space_before = Pt(26)
        inline(p, text.strip(), 11)
    doc.add_page_break()

    body = "# 1" + content.split("\n# 1", 1)[1]
    headings = re.findall(r"^# (.+)$", body, re.M)
    p = doc.add_paragraph("目录", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    for index, title in enumerate(headings, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        link(p, title, anchor=f"section_{index}", size=10.5)
    doc.add_page_break()

    lines = body.splitlines()
    i, hidx = 0, 0
    figures = 0
    last_heading = "系统结构"
    pending_caption = None
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("<!-- 图注"):
            pending_caption = line.removeprefix("<!-- 图注").removesuffix("-->").strip(" ：:") 
            i += 1
            continue
        if line.startswith("<!-- 截图"):
            name = line.removeprefix("<!-- 截图").removesuffix("-->").strip(" ：:")
            png = SCREENSHOTS / f"{name}.png"
            if not png.is_file():
                png = ASSETS / f"{name}.png"
            if not png.is_file():
                raise SystemExit(f"缺少配图 {name}.png（界面截图见 tools/capture_screenshots.mjs，"
                                 f"机械插图先跑 tools/blender/render_plates.py 再跑 tools/compose_figures.mjs）")
            figures += 1
            caption = pending_caption or name
            insert_figure(doc, png, f"图 {figures}  {caption}", caption)
            pending_caption = None
            i += 1
            continue
        if line.startswith("```mermaid"):
            block = []
            i += 1
            while i < len(lines) and lines[i].strip() != "```":
                block.append(lines[i])
                i += 1
            figures += 1
            png = render_mermaid("\n".join(block), f"figure-{figures}")
            caption = pending_caption or f"图 {figures}  {last_heading}"
            insert_figure(doc, png, f"图 {figures}  {caption}", f"LineGuard {caption}")
            pending_caption = None
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                row = [item.strip() for item in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", item) for item in row):
                    rows.append(row)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith("# "):
            title = line[2:]
            last_heading = title
            if title.startswith("附录"):
                doc.add_page_break()
            hidx += 1
            p = doc.add_paragraph(title, style="Heading 1")
            bookmark(p, hidx, f"section_{hidx}")
        elif line.startswith("## "):
            last_heading = line[3:]
            doc.add_paragraph(last_heading, style="Heading 2")
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if not re.match(r"^\d+\.", line) else WD_ALIGN_PARAGRAPH.LEFT
            if re.match(r"^\d+\.", line):
                p.paragraph_format.left_indent = Inches(0.15)
                p.paragraph_format.first_line_indent = Inches(-0.15)
            inline(p, line)
        i += 1
    # 清除内置标题样式的边框，避免 Word 默认蓝色横线残留。
    for root in (doc.styles.element, doc.element):
        for border_node in list(root.xpath(".//w:pBdr")):
            border_node.getparent().remove(border_node)
    doc.save(DEST)
    main_part = body.split("# 附录 A", 1)[0]
    main_han = len(re.findall(r"[\u4e00-\u9fff]", main_part))
    print(f"DOCX: {DEST}")
    print(f"Main Chinese characters: {main_han}")
    print(f"Sections: {len(headings)}; tables: {len(doc.tables)}; figures: {figures}")


if __name__ == "__main__":
    main()
