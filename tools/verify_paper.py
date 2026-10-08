"""从成品 DOCX 与实际渲染 PDF 验证课程格式，不以估算页数代替分页。"""
from pathlib import Path
from zipfile import ZipFile
import argparse
import json
import re
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.table import Table
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DEFAULT = ROOT / 'LineGuard_课程论文_V2.0.docx'

def font_data(run):
    fonts = run._element.rPr.rFonts if run._element.rPr is not None else None
    return (fonts.get(qn('w:eastAsia')) if fonts is not None else None,
            fonts.get(qn('w:ascii')) if fonts is not None else None,
            run.font.size.pt if run.font.size else None)

def indent(p):
    ind = p._p.pPr.find(qn('w:ind')) if p._p.pPr is not None else None
    return ind.get(qn('w:firstLineChars')) if ind is not None else None

def verify(docx, pdf):
    d = Document(docx)
    results = []
    def check(name, ok, detail=''):
        results.append({'rule': name, 'passed': bool(ok), 'detail': detail})
    check('A4 与四边 2.5 cm', all(abs(s.page_width.cm-21)<.02 and abs(s.page_height.cm-29.7)<.02
        and all(abs(getattr(s,a).cm-2.5)<.02 for a in ('top_margin','bottom_margin','left_margin','right_margin')) for s in d.sections))
    check('封面与正文分节',len(d.sections)==2)
    check('封面无页码', 'PAGE' not in d.sections[0].footer._element.xml)
    footer=d.sections[-1].footer
    check('正文页码居中从 1 起', 'PAGE' in footer._element.xml and not footer.is_linked_to_previous
        and footer.paragraphs[0].alignment == 1 and d.sections[-1]._sectPr.find(qn('w:pgNumType')).get(qn('w:start'))=='1')
    cover=d.tables[0]
    check('学号姓名填写正确',cover.cell(0,1).text=='2024303524' and cover.cell(1,1).text=='张立烁')
    check('题目已填写得分留空',bool(cover.cell(2,1).text) and not cover.cell(3,1).text)
    check('说明页已删除',not any('格式规范（本页删除）' in p.text for p in d.paragraphs))
    check('正文关闭隐式网格',d.sections[-1]._sectPr.find(qn('w:docGrid')) is None)
    body=[]; in_body=False
    for child in d.element.body:
        if child.tag==qn('w:p'):
            p=Paragraph(child,d)
            if p._p.pPr is not None and p._p.pPr.find(qn('w:sectPr')) is not None:
                in_body=True
                continue
            if in_body: body.append(p)
        elif child.tag==qn('w:tbl') and in_body: body.append(Table(child,d))
    paragraphs=[p for p in body if isinstance(p,Paragraph)]
    h1=[p for p in paragraphs if p.style.name=='Heading 1']
    check('一级标题另页',len(h1)==7 and all(p.paragraph_format.page_break_before for p in h1[1:]))
    check('二三级标题无强制分页',all(not p.paragraph_format.page_break_before for p in paragraphs if p.style.name in ('Heading 2','Heading 3')))
    for level,size in [(1,15),(2,14),(3,12)]:
        s=d.styles[f'Heading {level}']; rf=s._element.rPr.rFonts
        check(f'{level} 级标题黑体 {size} pt 加粗',s.font.bold and s.font.size.pt==size and rf.get(qn('w:eastAsia'))=='黑体')
    all_text=[p for p in d.paragraphs if p.text.strip()]
    all_text += [p for t in d.tables for row in t.rows for c in row.cells for p in c.paragraphs if p.text.strip()]
    check('全部文本 1.5 倍行距',all(p.paragraph_format.line_spacing==1.5 for p in all_text),str(len(all_text))+' 段（含表格与代码）')
    abstract=False; prose=[]; code=[]; refs=[]
    for p in paragraphs:
        if p.style.name=='Heading 1': abstract=p.text=='摘要'
        if not p.text.strip(): continue
        if p.style.name.startswith('Heading'):
            check('标题不首行缩进',indent(p) in (None,'0'),p.text)
            continue
        if p.runs and font_data(p.runs[0])[1]=='Consolas': code.append(p); continue
        if p.style.name in ('Caption','Code Caption'): continue
        if p.text.startswith('['): refs.append(p); continue
        if abstract:
            check('摘要不首行缩进',indent(p) in (None,'0'),p.text[:18])
        else: prose.append(p)
    check('正文宋体小四与两字符缩进',all(indent(p)=='200' and all(font_data(r)==('宋体','Times New Roman',12) for r in p.runs if r.text) for p in prose),str(len(prose))+' 段')
    check('代码 Consolas 五号中文宋体无首缩',len(code)>15 and all(indent(p) in (None,'0') and all(font_data(r)==('宋体','Consolas',10.5) for r in p.runs if r.text) for p in code))
    check('参考文献宋体五号无首缩',len(refs)==5 and all(indent(p) in (None,'0') and all(font_data(r)==('宋体','Times New Roman',10.5) for r in p.runs if r.text) for p in refs))
    captions=[p for p in paragraphs if p.style.name=='Caption']
    check('图表题宋体五号按章编号',len(captions)==5 and all(re.match(r'^[图表] \d+-\d+  ',p.text) and indent(p) in (None,'0') and all(font_data(r)==('宋体','Times New Roman',10.5) for r in p.runs if r.text) for p in captions))
    for i,element in enumerate(body):
        if isinstance(element,Table):
            check('表题在上方',i>0 and isinstance(body[i-1],Paragraph) and body[i-1].style.name=='Caption' and body[i-1].text.startswith('表 '))
            check('表格行不拆分且表头可重复',all(row._tr.trPr.find(qn('w:cantSplit')) is not None for row in element.rows) and element.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None)
        elif element._p.xpath('.//wp:extent'):
            check('图题紧随图片',i+1<len(body) and isinstance(body[i+1],Paragraph) and body[i+1].style.name=='Caption' and body[i+1].text.startswith('图 '))
            check('图片与图题保持同页',element.paragraph_format.keep_with_next)
    # 核心 Java 片段仅允许调整换行与缩进；必须能在真实源码中找到。
    joined=''.join(p.text for p in code)
    for start,end,file in [('public void check','if (commandRepository','TaskStateMachine.java'),('if (commandRepository','每个控制周期','MissionStateService.java')]:
        excerpt=joined.split(start,1)[1].split(end,1)[0]
        excerpt=start+excerpt
        source=(ROOT/'lineguard-system/backend/src/main/java/com/lineguard/service'/file).read_text()
        check('Java 节选与附件源码一致',re.sub(r'\s','',excerpt) in re.sub(r'\s','',source),file)
    reference=ROOT/'《智能机器人软件设计方法与实践》封面及格式规范.docx'
    with ZipFile(reference) as z1, ZipFile(docx) as z2:
        for name in ('word/media/image1.png','word/theme/theme1.xml','word/fontTable.xml','docProps/custom.xml'):
            check('模板保留部件一致',z1.read(name)==z2.read(name),name)
    check('实际 PDF 已生成',pdf.is_file(),str(pdf))
    if pdf.is_file():
        r=PdfReader(pdf); texts=[p.extract_text() for p in r.pages]
        check('实际总页数含封面在 5 到 10 页',5<=len(r.pages)<=10,str(len(r.pages)))
        check('PDF 中文与身份可提取',all(t in texts[0] for t in ('张立烁','2024303524','课程论文')))
        check('封面为一页', '2026-2027' in texts[0] and '摘要' in texts[1])
        check('正文页码连续',all(re.search(rf'(^|\n){i}\s*(\n|$)',t) for i,t in enumerate(texts[1:],1)))
        check('正文五章均已渲染',all(any(f'第 {i} 章' in t for t in texts) for i in range(1,6)))
    return results

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('docx',type=Path,nargs='?',default=DEFAULT)
    parser.add_argument('--pdf',type=Path)
    parser.add_argument('--json',type=Path)
    args=parser.parse_args()
    results=verify(args.docx,args.pdf or args.docx.with_suffix('.pdf'))
    for result in results:
        print(('PASS' if result['passed'] else 'FAIL'),result['rule'],result['detail'])
    if args.json:
        args.json.parent.mkdir(parents=True,exist_ok=True)
        args.json.write_text(json.dumps(results,ensure_ascii=False,indent=2))
    raise SystemExit(0 if all(r['passed'] for r in results) else 1)
